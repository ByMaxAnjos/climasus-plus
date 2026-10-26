#!/usr/bin/env python3
"""Gera a edição Word (.docx) a partir do HTML de impressão.

Fluxo: print-body.html -> docx.html (com capa e aberturas em imagem) ->
Pandoc com reference.docx -> pós-processamento python-docx.
"""
from __future__ import annotations

import re
import subprocess
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString, Tag
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent.parent
BOOK = ROOT / "_book"
REFERENCE = ROOT / "theme" / "reference.docx"

INK = RGBColor(0x0B, 0x24, 0x20)
TEAL = RGBColor(0x0E, 0x7C, 0x6B)
TEAL_DARK = RGBColor(0x0B, 0x5D, 0x52)
MUTED = RGBColor(0x5B, 0x6B, 0x64)
OCHRE = RGBColor(0xB7, 0x79, 0x1F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BODY = "Calibri"
HEAD = "Georgia"
MONO = "Consolas"


def make_docx_html() -> Path:
    soup = BeautifulSoup((BOOK / "print-body.html").read_text(encoding="utf-8"), "lxml")

    for tag in soup.find_all(["script", "style", "link", "title"]):
        tag.decompose()
    for meta in soup.find_all("meta", attrs={"name": "author"}):
        meta.decompose()

    # O kicker do sumário é redundante no Word (o cabeçalho já traz o título).
    for kicker in soup.select(".print-toc .subtitle"):
        kicker.decompose()

    # Separa kicker e título no sumário estático (o HTML usa spans sem espaço).
    for title in soup.select(".toc-part-title"):
        span = title.find("span")
        if span:
            span.insert_after(NavigableString(" — "))
    for item in soup.select(".print-toc li"):
        span = item.find("span", class_="toc-kicker")
        if span:
            span.insert_after(NavigableString(" — "))

    # Títulos de capítulo precisam ser Heading 1 no Word, não o estilo Title.
    for h1 in soup.select("h1.title"):
        del h1["class"]

    # Cartões da home não devem poluir o sumário/navegação do Word.
    for card in soup.select(".audience-card h3, .map-card h3"):
        p = soup.new_tag("p")
        strong = soup.new_tag("strong")
        strong.string = card.get_text(" ", strip=True)
        p.append(strong)
        card.replace_with(p)

    # Aberturas de parte: substitui o layout HTML pela arte em página inteira.
    parts = soup.select(".print-part")
    for i, part in enumerate(parts, start=1):
        img = soup.new_tag("img", src=f"images/part-openers/part-{i}-{['fundamentos','pipeline','gestor','pesquisador','casos'][i-1]}.jpg")
        img["alt"] = f"Parte {['I','II','III','IV','V'][i-1]}"
        wrapper = soup.new_tag("div", attrs={"class": "docx-part-image"})
        wrapper.append(img)
        part.replace_with(wrapper)

    # Capa.
    cover = soup.new_tag("div", attrs={"class": "docx-cover"})
    img = soup.new_tag("img", src="images/brand/capa-guia.jpg")
    img["alt"] = "Capa do guia prático do climasus+ Studio"
    cover.append(img)
    body = soup.find("body")
    body.insert(0, cover)

    out = BOOK / "docx.html"
    out.write_text(str(soup), encoding="utf-8")
    return out


def run_pandoc(html: Path) -> Path:
    out = BOOK / "docx-raw.docx"
    cmd = [
        "pandoc",
        html.name,
        "-f",
        "html",
        "-t",
        "docx",
        "--reference-doc",
        str(REFERENCE),
        "-o",
        out.name,
    ]
    # cwd = _book para que images/... resolvam ao lado do HTML.
    subprocess.run(cmd, check=True, cwd=BOOK)
    return out


def set_run_font(run, name: str, size: float, color: RGBColor, bold: bool | None = None, italic: bool | None = None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def cell_margins(table, top=40, bottom=40, left=80, right=80):
    tbl_pr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tbl_pr.append(mar)


def table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "D8DED6")
        borders.append(el)
    tbl_pr.append(borders)


def set_table_width(table):
    tbl_pr = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")
    tbl_pr.append(width)


def style_tables(doc: Document):
    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_table_width(table)
        cell_margins(table)
        table_borders(table)

        if not table.rows:
            continue

        first = table.rows[0]
        empty_header = all(not c.text.strip() for c in first.cells)
        if empty_header:
            table._tbl.remove(first._tr)
            header_index = None
        else:
            header_index = 0

        for ri, row in enumerate(table.rows):
            if header_index is not None and ri == header_index:
                for cell in row.cells:
                    shade_cell(cell, "0B2420")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            set_run_font(run, BODY, 7.8, WHITE, bold=True)
                            run.font.all_caps = True
            else:
                for cell in row.cells:
                    if ri % 2 == 0:
                        shade_cell(cell, "F2F6F1")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            if run.font.name and "Mono" in (run.font.name or ""):
                                set_run_font(run, MONO, 7.2, TEAL_DARK)
                            else:
                                set_run_font(run, BODY, 8.2, RGBColor(0x16, 0x3B, 0x34))


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def add_headers_footers(doc: Document):
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.text = ""
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.45), WD_TAB_ALIGNMENT.RIGHT)
    r = p.add_run("climasus+ Studio · Guia prático")
    set_run_font(r, BODY, 8, MUTED)
    r = p.add_run("\tApoio: INCT-CONEXAO")
    set_run_font(r, BODY, 8, OCHRE, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Apoio: INCT-CONEXAO · página ")
    set_run_font(r, BODY, 8, MUTED)
    add_page_field(p)
    for run in p.runs:
        set_run_font(run, BODY, 8, MUTED)


def resize_images(doc: Document):
    """Ajusta todos os wp:inline, inclusive imagens que python-docx não lista."""
    EMU = 914400
    logos = {"climasus+ Studio", "INCT-CONEXAO", "Fiocruz Rondônia", "CNPq", "UFJF"}

    for inline in doc.element.body.iter(qn("wp:inline")):
        extent = inline.find(qn("wp:extent"))
        doc_pr = inline.find(qn("wp:docPr"))
        if extent is None:
            continue
        descr = (doc_pr.get("descr") if doc_pr is not None else "") or ""
        old_cx = int(extent.get("cx", "0"))
        old_cy = int(extent.get("cy", "0"))
        if old_cx <= 0 or old_cy <= 0:
            continue

        p_el = inline.xpath("ancestor::w:p")[0]
        para = Paragraph(p_el, doc._body)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(6)

        if descr.startswith("Capa"):
            target_w = 6.3
            target_h = 9.45
            para.paragraph_format.space_after = Pt(0)
        elif descr.startswith("Parte"):
            target_w = 6.3
            target_h = 8.91
            para.paragraph_format.page_break_before = True
            para.paragraph_format.space_after = Pt(0)
        elif descr in logos:
            target_w = 0.7 if descr == "INCT-CONEXAO" else 0.55
            target_h = target_w * old_cy / old_cx
        else:
            target_w = 6.3
            target_h = target_w * old_cy / old_cx
            if target_h > 8.5:
                target_h = 8.5
                target_w = target_h * old_cx / old_cy

        extent.set("cx", str(int(target_w * EMU)))
        extent.set("cy", str(int(target_h * EMU)))


def style_paragraphs(doc: Document):
    for p in doc.paragraphs:
        text = p.text.strip()
        if not text:
            continue

        # Sumário estático: parte e capítulo em duas vozes tipográficas.
        if text.startswith("Parte ") and " — " in text:
            kicker, title = text.split(" — ", 1)
            p.clear()
            r1 = p.add_run(kicker)
            set_run_font(r1, BODY, 8, OCHRE, bold=True)
            r1.font.all_caps = True
            r1.add_break()
            r2 = p.add_run(title)
            set_run_font(r2, HEAD, 12.5, INK, bold=True)
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            continue

        if text.startswith("cap. ") and " — " in text:
            kicker, title = text.split(" — ", 1)
            p.clear()
            r1 = p.add_run(kicker)
            set_run_font(r1, BODY, 7.5, OCHRE, bold=True)
            r1.font.all_caps = True
            r2 = p.add_run(" — " + title)
            set_run_font(r2, BODY, 9.3, RGBColor(0x16, 0x3B, 0x34))
            p.paragraph_format.space_after = Pt(2)
            continue

        # Kicker de capítulo / seção.
        if re.match(r"^Capítulo\s+\d+", text) or text in {"Ficha técnica e créditos", "climasus+ Studio · guia prático"}:
            p.style = doc.styles["Normal"]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            if text == "climasus+ Studio · guia prático":
                p.paragraph_format.page_break_before = True
                p.paragraph_format.space_before = Pt(0)
            for run in p.runs:
                set_run_font(run, BODY, 8, OCHRE, bold=True)
                run.font.all_caps = True
            continue

        if "Verifique se ficou" in text:
            p.style = doc.styles["Heading 3"]
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            for run in p.runs:
                set_run_font(run, HEAD, 12, TEAL, bold=True)
            continue

        if "Como ler as tabelas" in text:
            p.style = doc.styles["Normal"]
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            for run in p.runs:
                set_run_font(run, BODY, 9.5, TEAL, bold=True)
            continue

        # Título da caixa de destaque.
        if text in {"Dica", "Nota", "Importante", "Atenção", "Cuidado"}:
            p.style = doc.styles["Heading 4"]
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            for run in p.runs:
                set_run_font(run, BODY, 9, TEAL, bold=True, italic=False)
                run.font.all_caps = False
            continue

        # Cada H1 começa em página nova (capa e aberturas usam imagens).
        if p.style.name == "Heading 1":
            p.paragraph_format.page_break_before = True

        # Primeira seção da apresentação começa em página nova.
        if text == "Para quem é este guia":
            p.paragraph_format.page_break_before = True


def force_heading_fonts(doc: Document):
    specs = {
        "Heading 1": (HEAD, 23, INK, True),
        "Heading 2": (HEAD, 15, TEAL, True),
        "Heading 3": (BODY, 11, TEAL_DARK, True),
        "Heading 4": (BODY, 8.5, OCHRE, True),
    }
    for p in doc.paragraphs:
        if p.style.name in specs:
            name, size, color, bold = specs[p.style.name]
            if p.style.name == "Heading 3" and "Verifique se ficou" in p.text:
                name, size, color, bold = HEAD, 12, TEAL, True
            for run in p.runs:
                set_run_font(run, name, size, color, bold=bold)


def build():
    if not REFERENCE.exists():
        subprocess.run(["python3", str(ROOT / "scripts" / "make-reference-docx.py")], check=True)

    html = make_docx_html()
    raw = run_pandoc(html)
    doc = Document(raw)

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    # Fundo de página em tom de papel.
    bg = OxmlElement("w:background")
    bg.set(qn("w:color"), "F7F4EC")
    doc.element.insert(0, bg)

    resize_images(doc)
    style_paragraphs(doc)
    force_heading_fonts(doc)
    style_tables(doc)
    add_headers_footers(doc)

    doc.core_properties.title = "climasus+ Studio · Guia prático sobre clima, saúde e Amazônia"
    doc.core_properties.author = "Max Anjos · ecossistema climaSUS"
    doc.core_properties.subject = "Clima, saúde e Amazônia"
    doc.core_properties.keywords = "climasus, clima, saúde, Amazônia, DATASUS, SUS, Fiocruz, INCT-CONEXAO"
    doc.core_properties.comments = "Edição Word gerada a partir do guia prático do climasus+ Studio."

    out = BOOK / "climasus-studio-guia-pratico.docx"
    doc.save(out)
    print(f"DOCX gerado em {out} ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    build()
